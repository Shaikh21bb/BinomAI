import csv
import io
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from statistics import median
from typing import Any, Iterable, Optional

from app.db.models.product_search import ProductSearchItem
from app.db.models.project import Project
from app.db.models.sourcing import SupplierOffer


ZERO = Decimal("0")
ONE = Decimal("1")

UNIT_ALIASES = {
    "шт": "pcs", "штук": "pcs", "штука": "pcs", "pcs": "pcs", "piece": "pcs",
    "компл": "set", "комплект": "set", "набор": "set", "set": "set",
    "кг": "kg", "kg": "kg", "килограмм": "kg", "килограммов": "kg",
    "г": "g", "гр": "g", "g": "g",
    "т": "t", "тн": "t", "тонна": "t", "тонн": "t", "ton": "t",
    "м": "m", "пм": "m", "погм": "m", "погонныйм": "m", "meter": "m",
    "см": "cm", "cm": "cm", "мм": "mm", "mm": "mm",
    "м2": "m2", "м²": "m2", "sqm": "m2",
    "см2": "cm2", "см²": "cm2", "cm2": "cm2",
    "м3": "m3", "м³": "m3", "cbm": "m3",
    "л": "l", "литр": "l", "литров": "l", "liter": "l",
}

UNIT_META = {
    "pcs": ("count", Decimal("1")),
    "set": ("set", Decimal("1")),
    "kg": ("mass", Decimal("1")),
    "g": ("mass", Decimal("0.001")),
    "t": ("mass", Decimal("1000")),
    "m": ("length", Decimal("1")),
    "cm": ("length", Decimal("0.01")),
    "mm": ("length", Decimal("0.001")),
    "m2": ("area", Decimal("1")),
    "cm2": ("area", Decimal("0.0001")),
    "m3": ("volume", Decimal("1")),
    "l": ("volume", Decimal("0.001")),
}

HEADER_ALIASES = {
    "supplier_name": {"supplier", "suppliername", "поставщик", "наименованиепоставщика"},
    "supplier_bin": {"bin", "supplierbin", "бин", "бинпоставщика"},
    "supplier_contact": {"contact", "contacts", "контакт", "контакты"},
    "item_id": {"itemid", "lineitemid", "позицияid", "idпозиции"},
    "original_item_name": {"item", "itemname", "product", "productname", "товар", "наименование", "позиция"},
    "original_unit": {"unit", "единица", "едизм", "единицаизмерения"},
    "quoted_quantity": {"quantity", "qty", "количество", "колво"},
    "unit_price": {"unitprice", "price", "цена", "ценазаединицу"},
    "price_quantity": {"pricequantity", "priceper", "ценаза", "кратностьцены"},
    "currency": {"currency", "валюта"},
    "exchange_rate_to_kzt": {"exchangerate", "ratetokzt", "курсктенге", "курс"},
    "vat_included": {"vatincluded", "ндсвключен", "сндс"},
    "vat_rate": {"vatrate", "ндсставка", "ндс"},
    "moq": {"moq", "minorder", "минимальныйзаказ", "минзаказ"},
    "available_quantity": {"available", "stock", "наличие", "доступноколичество"},
    "delivery_cost": {"deliverycost", "доставка", "стоимостьдоставки"},
    "lead_time_days": {"leadtime", "leadtimedays", "срокдоставкидней", "днейдоставки"},
    "warranty_months": {"warranty", "warrantymonths", "гарантиямесяцев", "гарантия"},
    "certificates": {"certificates", "сертификаты"},
    "compliance_status": {"compliance", "соответствие", "статуссоответствия"},
    "compliance_notes": {"notes", "примечание", "комментарий"},
    "quote_date": {"quotedate", "датапредложения", "датакп"},
    "valid_until": {"validuntil", "действительнодо", "срокдействия"},
    "source_url": {"url", "link", "ссылка"},
}


def normalize_name(value: str) -> str:
    text = (value or "").lower().replace("ё", "е")
    text = re.sub(r"[^a-zа-я0-9]+", " ", text)
    return " ".join(text.split())


def normalize_unit(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    key = value.lower().strip().replace(".", "").replace(" ", "")
    return UNIT_ALIASES.get(key)


def conversion_factor(from_unit: Optional[str], to_unit: Optional[str]) -> Optional[Decimal]:
    """Return the safe multiplier from one unit to another, or None if ambiguous."""
    source = normalize_unit(from_unit)
    target = normalize_unit(to_unit)
    if not source or not target:
        return None
    source_meta = UNIT_META[source]
    target_meta = UNIT_META[target]
    if source_meta[0] != target_meta[0]:
        return None
    return source_meta[1] / target_meta[1]


def match_line_item(name: str, items: Iterable[ProductSearchItem]) -> tuple[Optional[ProductSearchItem], str, Decimal]:
    needle = normalize_name(name)
    best_item: Optional[ProductSearchItem] = None
    best_score = 0.0
    second_score = 0.0
    needle_tokens = set(needle.split())
    for item in items:
        candidate = normalize_name(item.product_name)
        candidate_tokens = set(candidate.split())
        sequence = SequenceMatcher(None, needle, candidate).ratio()
        union = needle_tokens | candidate_tokens
        overlap = len(needle_tokens & candidate_tokens) / len(union) if union else 0.0
        score = sequence * 0.65 + overlap * 0.35
        if score > best_score:
            second_score = best_score
            best_score = score
            best_item = item
        elif score > second_score:
            second_score = score

    confidence = Decimal(str(round(best_score, 4)))
    if best_item and best_score >= 0.86 and best_score - second_score >= 0.08:
        return best_item, "matched", confidence
    if best_item and best_score >= 0.58:
        return best_item, "needs_review", confidence
    return None, "unmatched", confidence


def _decimal(value: Any, default: Optional[Decimal] = None) -> Optional[Decimal]:
    if value is None or value == "":
        return default
    if isinstance(value, Decimal):
        return value
    try:
        cleaned = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
        return Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return default


def _round_money(value: Optional[Decimal]) -> Optional[float]:
    if value is None:
        return None
    return float(value.quantize(Decimal("0.01")))


def _flag(code: str, level: str, message: str) -> dict[str, str]:
    return {"code": code, "level": level, "message": message}


def evaluate_offer(
    offer: SupplierOffer,
    item: ProductSearchItem,
    project: Project,
    today: Optional[date] = None,
) -> dict[str, Any]:
    today = today or datetime.now(timezone.utc).date()
    flags: list[dict[str, str]] = []
    reasons: list[str] = []
    required = _decimal(item.quantity)
    factor = conversion_factor(item.unit, offer.original_unit)

    if required is None or required <= ZERO:
        flags.append(_flag("missing_quantity", "error", "В позиции не указано корректное количество"))
        required_offer_qty = None
    elif factor is None:
        required_offer_qty = None
        flags.append(_flag("unsafe_unit_conversion", "error", "Единицы нельзя безопасно сопоставить без проверки"))
    else:
        required_offer_qty = required * factor

    rate = ONE if offer.currency == "KZT" else _decimal(offer.exchange_rate_to_kzt)
    if rate is None:
        flags.append(_flag("missing_exchange_rate", "error", f"Для {offer.currency} не указан курс к KZT"))

    price_quantity = _decimal(offer.price_quantity, ONE) or ONE
    if price_quantity <= ZERO:
        flags.append(_flag("invalid_price_quantity", "error", "Кратность цены должна быть больше нуля"))

    vat_multiplier: Optional[Decimal]
    if offer.vat_included is True:
        vat_multiplier = ONE
    elif offer.vat_included is False and offer.vat_rate is not None:
        vat_multiplier = ONE + (_decimal(offer.vat_rate, ZERO) or ZERO) / Decimal("100")
    else:
        vat_multiplier = None
        flags.append(_flag("unknown_vat", "warning", "Неясно, включён ли НДС; итоговая стоимость не подтверждена"))

    purchased_qty = required_offer_qty
    moq = _decimal(offer.moq)
    if purchased_qty is not None and moq and moq > ZERO:
        purchased_qty = max(purchased_qty, moq)
        if purchased_qty > required_offer_qty:
            reasons.append(f"Учтён минимальный заказ: {purchased_qty:g} {offer.original_unit or ''}".strip())

    available = _decimal(offer.available_quantity)
    if purchased_qty is not None and available is not None and available < purchased_qty:
        flags.append(_flag("insufficient_availability", "error", "Подтверждённого наличия недостаточно для требуемого объёма"))
    elif available is None:
        flags.append(_flag("unknown_availability", "warning", "Наличие не подтверждено"))

    if offer.valid_until and offer.valid_until < today:
        flags.append(_flag("expired_quote", "error", "Срок действия предложения истёк"))
    elif offer.quote_date is None:
        flags.append(_flag("unknown_price_date", "warning", "Дата цены не указана"))
    elif (today - offer.quote_date).days > 30:
        flags.append(_flag("stale_price", "warning", "Цена старше 30 дней"))

    deadline = project.deadline_at.date() if project.deadline_at else None
    if deadline and offer.lead_time_days is not None:
        arrival = today + timedelta(days=offer.lead_time_days)
        if arrival > deadline:
            flags.append(_flag("deadline_risk", "error", f"Поставка ожидается после срока тендера ({deadline:%d.%m.%Y})"))
        elif arrival + timedelta(days=3) > deadline:
            flags.append(_flag("deadline_buffer", "warning", "До срока тендера остаётся менее трёх дней запаса"))
    elif offer.lead_time_days is None:
        flags.append(_flag("unknown_lead_time", "warning", "Срок поставки не указан"))

    if not offer.certificates:
        flags.append(_flag("missing_certificates", "warning", "Сертификаты не приложены"))
    if offer.warranty_months is None:
        flags.append(_flag("missing_warranty", "warning", "Условия гарантии не указаны"))
    if not offer.characteristics:
        flags.append(_flag("missing_characteristics", "warning", "Характеристики предложения не заполнены"))

    if offer.compliance_status == "noncompliant":
        flags.append(_flag("noncompliant", "error", "Предложение отмечено как несоответствующее"))
    elif offer.compliance_status == "partial":
        flags.append(_flag("partial_compliance", "warning", "Соответствие подтверждено частично"))
    elif offer.compliance_status == "unknown":
        flags.append(_flag("unknown_compliance", "error", "Соответствие техническому заданию не проверено"))

    landed: Optional[Decimal] = None
    if purchased_qty is not None and rate is not None and vat_multiplier is not None and price_quantity > ZERO:
        goods = purchased_qty * _decimal(offer.unit_price, ZERO) / price_quantity
        delivery = _decimal(offer.delivery_cost, ZERO) or ZERO
        landed = (goods + delivery) * rate * vat_multiplier

    error_codes = {flag["code"] for flag in flags if flag["level"] == "error"}
    eligible = landed is not None and not error_codes
    if offer.compliance_status == "compliant":
        reasons.append("Соответствие ТЗ подтверждено")
    if offer.certificates:
        reasons.append(f"Указаны сертификаты: {len(offer.certificates)}")
    if offer.warranty_months:
        reasons.append(f"Гарантия {offer.warranty_months} мес.")
    if offer.lead_time_days is not None:
        reasons.append(f"Срок поставки {offer.lead_time_days} дн.")

    return {
        "offer": _offer_dict(offer),
        "required_offer_quantity": float(required_offer_qty) if required_offer_qty is not None else None,
        "purchased_quantity": float(purchased_qty) if purchased_qty is not None else None,
        "landed_cost_kzt": _round_money(landed),
        "landed_unit_cost_kzt": _round_money(landed / required) if landed is not None and required else None,
        "score": 0.0,
        "eligible": eligible,
        "flags": flags,
        "reasons": reasons,
    }


def _offer_dict(offer: SupplierOffer) -> dict[str, Any]:
    fields = (
        "id", "project_id", "company_id", "item_id", "supplier_name", "supplier_bin",
        "supplier_contact", "original_item_name", "normalized_item_name", "original_unit",
        "normalized_unit", "quoted_quantity", "unit_price", "price_quantity", "currency",
        "exchange_rate_to_kzt", "vat_included", "vat_rate", "moq", "available_quantity",
        "delivery_cost", "lead_time_days", "warranty_months", "certificates", "characteristics",
        "compliance_status", "compliance_notes", "quote_date", "valid_until", "source_type",
        "source_filename", "source_url", "match_status", "match_confidence", "is_selected",
        "selection_note", "created_at", "updated_at",
    )
    return {field: getattr(offer, field, None) for field in fields}


def build_comparison(
    project: Project,
    items: list[ProductSearchItem],
    offers: list[SupplierOffer],
    target_margin_pct: Decimal,
    today: Optional[date] = None,
) -> dict[str, Any]:
    today = today or datetime.now(timezone.utc).date()
    offers_by_item: dict[Any, list[SupplierOffer]] = {}
    unmatched = []
    for offer in offers:
        if offer.item_id is None or offer.match_status != "matched":
            unmatched.append(_offer_dict(offer))
        else:
            offers_by_item.setdefault(offer.item_id, []).append(offer)

    item_rows = []
    chosen_total = ZERO
    covered_items = 0
    for item in items:
        evaluated = [evaluate_offer(offer, item, project, today) for offer in offers_by_item.get(item.id, [])]
        eligible_costs = [row["landed_cost_kzt"] for row in evaluated if row["eligible"] and row["landed_cost_kzt"] is not None]
        cost_median = median(eligible_costs) if eligible_costs else None
        min_cost = min(eligible_costs) if eligible_costs else None

        for row in evaluated:
            cost = row["landed_cost_kzt"]
            offer_data = row["offer"]
            if cost_median is not None and len(eligible_costs) >= 3 and cost < cost_median * 0.65:
                row["flags"].append(_flag("suspiciously_cheap", "warning", "Цена более чем на 35% ниже медианы; проверьте состав и условия"))

            compliance_points = {"compliant": 45, "partial": 25, "unknown": 5, "noncompliant": 0}.get(offer_data["compliance_status"], 0)
            economics_points = 30 * (min_cost / cost) if min_cost and cost else 0
            lead = offer_data["lead_time_days"]
            delivery_points = 3 if lead is None else 15 if lead <= 7 else 12 if lead <= 14 else 8 if lead <= 30 else 4
            assurance_points = (4 if offer_data["certificates"] else 0) + (3 if offer_data["warranty_months"] is not None else 0)
            quote_date = offer_data["quote_date"]
            if quote_date and (today - quote_date).days <= 30:
                assurance_points += 3
            row["score"] = round(float(compliance_points + economics_points + delivery_points + assurance_points), 1)
            if min_cost and cost and cost == min_cost:
                row["reasons"].append("Минимальная подтверждённая приведённая стоимость среди допустимых вариантов")

        evaluated.sort(key=lambda row: (not row["eligible"], -row["score"], row["landed_cost_kzt"] or float("inf")))
        recommended = next((row for row in evaluated if row["eligible"]), None)
        selected = next((row for row in evaluated if row["offer"]["is_selected"]), None) or recommended
        if selected and selected["landed_cost_kzt"] is not None:
            chosen_total += Decimal(str(selected["landed_cost_kzt"]))
            covered_items += 1

        item_rows.append({
            "item": {
                "id": item.id,
                "product_name": item.product_name,
                "specs": item.specs,
                "unit": item.unit,
                "normalized_unit": normalize_unit(item.unit),
                "quantity": item.quantity,
                "source_section": item.source_section,
                "status": item.status,
                "discovery_leads": item.results or [],
            },
            "offers": evaluated,
            "recommended_offer_id": recommended["offer"]["id"] if recommended else None,
            "selected_offer_id": selected["offer"]["id"] if selected else None,
            "selection_is_manual": bool(selected and selected["offer"]["is_selected"]),
        })

    margin = _decimal(target_margin_pct, Decimal("15")) or Decimal("15")
    estimated_bid = chosen_total / (ONE - margin / Decimal("100")) if covered_items and margin < 100 else ZERO
    incomplete = max(0, len(items) - covered_items)
    return {
        "project": {
            "id": project.id,
            "name": project.name,
            "deadline_at": project.deadline_at,
            "customer_name": project.customer_name,
        },
        "settings": {"target_margin_pct": margin, "base_currency": "KZT"},
        "items": item_rows,
        "unmatched_offers": unmatched,
        "summary": {
            "line_items": len(items),
            "covered_items": covered_items,
            "incomplete_items": incomplete,
            "estimated_cost_kzt": _round_money(chosen_total),
            "estimated_bid_kzt": _round_money(estimated_bid) if covered_items else None,
            "estimated_profit_kzt": _round_money(estimated_bid - chosen_total) if covered_items else None,
            "is_complete": incomplete == 0 and len(items) > 0,
            "caveat": "Расчёт является оценкой и требует проверки условий КП, налогов, наличия и логистики." + (" Не все позиции покрыты." if incomplete else ""),
        },
    }


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-zа-я0-9]", "", str(value or "").lower().replace("ё", "е"))


def _canonicalize_row(row: dict[str, Any]) -> dict[str, Any]:
    alias_lookup = {alias: canonical for canonical, aliases in HEADER_ALIASES.items() for alias in aliases}
    return {alias_lookup.get(_header_key(key), _header_key(key)): value for key, value in row.items()}


def parse_quote_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    elif suffix == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX import requires the openpyxl package") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            return []
        rows = [dict(zip(headers, values_row)) for values_row in values]
    else:
        raise ValueError("Поддерживаются только CSV и XLSX")
    if len(rows) > 1000:
        raise ValueError("Один файл может содержать не более 1000 строк")
    return [_canonicalize_row(row) for row in rows if any(value not in (None, "") for value in row.values())]


def parse_bool(value: Any) -> Optional[bool]:
    if value is None or value == "":
        return None
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "да", "сндс", "включен", "включено"}:
        return True
    if text in {"0", "false", "no", "n", "нет", "безндс", "невключен", "невключено"}:
        return False
    return None


def parse_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def quote_row_payload(row: dict[str, Any]) -> dict[str, Any]:
    compliance_aliases = {
        "соответствует": "compliant", "compliant": "compliant", "да": "compliant",
        "частично": "partial", "partial": "partial",
        "не соответствует": "noncompliant", "несоответствует": "noncompliant", "noncompliant": "noncompliant", "нет": "noncompliant",
        "неизвестно": "unknown", "unknown": "unknown", "": "unknown",
    }
    certificates = row.get("certificates")
    certificate_list = [part.strip() for part in re.split(r"[;,|]", str(certificates)) if part.strip()] if certificates else []
    compliance = compliance_aliases.get(normalize_name(str(row.get("compliance_status") or "")), "unknown")
    currency = str(row.get("currency") or "KZT").strip().upper()
    currency = {"ТГ": "KZT", "ТЕНГЕ": "KZT", "₸": "KZT", "$": "USD", "€": "EUR"}.get(currency, currency)
    return {
        "supplier_name": str(row.get("supplier_name") or "").strip(),
        "supplier_bin": str(row.get("supplier_bin") or "").strip() or None,
        "supplier_contact": str(row.get("supplier_contact") or "").strip() or None,
        "original_item_name": str(row.get("original_item_name") or "").strip(),
        "original_unit": str(row.get("original_unit") or "").strip() or None,
        "quoted_quantity": _decimal(row.get("quoted_quantity")),
        "unit_price": _decimal(row.get("unit_price")),
        "price_quantity": _decimal(row.get("price_quantity"), ONE),
        "currency": currency,
        "exchange_rate_to_kzt": _decimal(row.get("exchange_rate_to_kzt")),
        "vat_included": parse_bool(row.get("vat_included")),
        "vat_rate": _decimal(row.get("vat_rate")),
        "moq": _decimal(row.get("moq")),
        "available_quantity": _decimal(row.get("available_quantity")),
        "delivery_cost": _decimal(row.get("delivery_cost")),
        "lead_time_days": int(_decimal(row.get("lead_time_days"))) if _decimal(row.get("lead_time_days")) is not None else None,
        "warranty_months": int(_decimal(row.get("warranty_months"))) if _decimal(row.get("warranty_months")) is not None else None,
        "certificates": certificate_list,
        "characteristics": {},
        "compliance_status": compliance,
        "compliance_notes": str(row.get("compliance_notes") or "").strip() or None,
        "quote_date": parse_date(row.get("quote_date")),
        "valid_until": parse_date(row.get("valid_until")),
        "source_url": str(row.get("source_url") or "").strip() or None,
    }


def build_rfq_draft(
    project: Project,
    items: list[ProductSearchItem],
    response_deadline: Optional[date],
    delivery_location: Optional[str],
    notes: Optional[str],
) -> tuple[str, str]:
    subject = f"Запрос коммерческого предложения — {project.name}"
    lines = [
        "Добрый день!",
        "",
        f"Просим предоставить коммерческое предложение для участия в проекте «{project.name}».",
        "",
        "Перечень позиций:",
    ]
    for index, item in enumerate(items, 1):
        qty = f" — {item.quantity:g} {item.unit or ''}" if item.quantity is not None else ""
        specs = f"; характеристики: {item.specs}" if item.specs else ""
        lines.append(f"{index}. {item.product_name}{qty}{specs}")
    lines.extend([
        "",
        "В КП просим отдельно указать цену за единицу и итоговую сумму, валюту, НДС, минимальный объём заказа, наличие, стоимость и срок доставки, гарантию, сертификаты и срок действия предложения.",
    ])
    if delivery_location:
        lines.append(f"Место поставки: {delivery_location}.")
    if response_deadline:
        lines.append(f"Просим направить ответ до {response_deadline:%d.%m.%Y}.")
    if notes:
        lines.extend(["", f"Дополнительные условия: {notes}"])
    lines.extend(["", "С уважением,"])
    return subject, "\n".join(lines)
